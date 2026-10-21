"""Entregables: memoria de cálculo, hoja de datos e informe técnico.

Cada documento se arma una sola vez como una lista de bloques
(títulos, párrafos, tablas, gráfico) y se presenta de dos formas:
HTML (vista previa en el programa e impresión) y Word editable (.docx).
Todos los valores se muestran en las unidades de presentación elegidas.
"""
from __future__ import annotations

import datetime as _dt
import html as _html
import io
import math
from dataclasses import dataclass, field

from . import __version__, compresible as CG, geometria as G, tuberias, unidades as U
from .modelos import Severidad
from .proyecto import FUGA_CIERRE, REF_CIERRE, Resultado

EMPRESA = "INSTRUMATIK S.A.S."
PRODUCTO = "N25 Valves Sizing"
# Texto de marca de agua de la licencia de demostración (None con licencia comercial)
MARCA_AGUA: str | None = None
NOMBRE_SEV = {Severidad.INFO: "Información", Severidad.ADVERTENCIA: "Advertencia",
              Severidad.CRITICA: "Crítica", Severidad.ERROR: "Error de datos"}
COLOR_SEV = {Severidad.INFO: "2F6FB0", Severidad.ADVERTENCIA: "B07A00",
             Severidad.CRITICA: "C0392B", Severidad.ERROR: "7D1F9E"}
CARACTERISTICA = {"isoporcentual": "Isoporcentual", "lineal": "Lineal", "apertura_rapida": "Apertura rápida"}


def fmt(v, cifras: int = 4) -> str:
    """Número con cifras significativas y formato español; «—» si no hay valor."""
    if v is None or (isinstance(v, float) and not math.isfinite(v)):
        return "—"
    if isinstance(v, str):
        return v
    if v == 0:
        return "0"
    a = abs(v)
    if a >= 1e7 or a < 1e-4:
        return f"{v:.{cifras - 1}e}".replace(".", ",")
    dec = max(0, cifras - 1 - int(math.floor(math.log10(a))))
    return f"{v:,.{dec}f}".replace(",", "X").replace(".", ",").replace("X", ".")


@dataclass
class Presentacion:
    """Unidad de presentación por magnitud (las mismas que el usuario eligió en las entradas)."""
    u: dict = field(default_factory=lambda: dict(U.SISTEMAS["SI"]))

    def un(self, magnitud: str) -> str:
        if magnitud == "presion_abs" and magnitud not in self.u:
            return {"bar g": "bar a", "psig": "psia", "kPa g": "kPa a", "MPa g": "MPa a",
                    "kg/cm² g": "kg/cm² a"}.get(self.u["presion"], self.u["presion"])
        return self.u[magnitud]

    def v(self, valor, magnitud: str, cifras: int = 4) -> str:
        if valor is None:
            return "—"
        base = "presion" if magnitud == "presion_abs" else magnitud
        return fmt(U.desde_base(valor, base, self.un(magnitud)), cifras)

    def vu(self, valor, magnitud: str, cifras: int = 4) -> str:
        return f"{self.v(valor, magnitud, cifras)} {self.un(magnitud)}"


# ---------------------------------------------------------------------------
# Modelo de documento
# ---------------------------------------------------------------------------
@dataclass
class Bloque:
    tipo: str            # h1, h2, h3, p, nota, lista, tabla, grafico, alerta, salto
    contenido: object = None
    opciones: dict = field(default_factory=dict)


@dataclass
class Documento:
    titulo: str
    codigo: str          # sufijo del nombre de archivo
    res: Resultado
    bloques: list[Bloque] = field(default_factory=list)

    def h1(self, t): self.bloques.append(Bloque("h1", t))
    def h2(self, t): self.bloques.append(Bloque("h2", t))
    def h3(self, t): self.bloques.append(Bloque("h3", t))
    def p(self, t): self.bloques.append(Bloque("p", t))
    def nota(self, t): self.bloques.append(Bloque("nota", t))
    def lista(self, items): self.bloques.append(Bloque("lista", list(items)))
    def grafico(self): self.bloques.append(Bloque("grafico"))
    def alerta(self, caso, a): self.bloques.append(Bloque("alerta", (caso, a)))

    def tabla(self, filas, encabezado=None, anchos=None, resaltar_col=None):
        """filas: lista de listas de texto. resaltar_col: índice desde el cual las columnas son resultados."""
        self.bloques.append(Bloque("tabla", filas, {"encabezado": encabezado, "anchos": anchos,
                                                    "resaltar_col": resaltar_col}))


# ---------------------------------------------------------------------------
# Contenido común
# ---------------------------------------------------------------------------
def _casos(res):
    return [rc.caso.nombre for rc in res.informe.casos]


def _identificacion(doc: Documento):
    i = doc.res.ident
    fecha = i.fecha or _dt.date.today().strftime("%d/%m/%Y")
    doc.tabla([["Proyecto", i.proyecto or "—", "Tag", i.tag or "—"],
               ["Cliente", i.cliente or "—", "Servicio", i.servicio or "—"],
               ["Documento", i.documento or "—", "Revisión", i.revision or "—"],
               ["Preparado por", i.preparado or "—", "Fecha", fecha]], anchos=[18, 32, 18, 32])


def _fila(et, unidad, valores, pr: Presentacion | None = None, magnitud=None, cifras=4):
    if magnitud:
        vals = [pr.v(v, magnitud, cifras) for v in valores]
        unidad = pr.un(magnitud)
    else:
        vals = [fmt(v, cifras) if not isinstance(v, str) else v for v in valores]
    return [et, unidad, *vals]


def _tabla_entradas(res: Resultado, pr: Presentacion):
    d = res.dim
    filas = []
    W = [c.W for c in d.casos]
    if all(c.Q is not None for c in d.casos):
        filas.append(_fila("Caudal volumétrico", "", [c.Q for c in d.casos], pr, "caudal_vol"))
    else:
        filas.append(_fila("Caudal másico", "", W, pr, "caudal_masico"))
    filas += [
        _fila("Presión de entrada P1", "", [c.P1 for c in d.casos], pr, "presion"),
        _fila("Presión de salida P2", "", [c.P2 for c in d.casos], pr, "presion"),
        _fila("Temperatura de entrada T1", "", [c.T1 for c in d.casos], pr, "temperatura"),
        _fila("Gravedad específica SG", "-", [c.SG for c in d.casos]),
        _fila("Viscosidad dinámica μ", "", [c.mu for c in d.casos], pr, "viscosidad_din"),
        _fila("Presión de vapor Pv", "", [c.Pv for c in d.casos], pr, "presion_abs"),
        _fila("Factor de recuperación FL", "-", [c.FL for c in d.casos], cifras=3),
        _fila("Modificador de estilo Fd", "-", [c.Fd for c in d.casos], cifras=3),
        _fila("Coef. de cavitación del trim Kc", "-",
              [f"{fmt(x['Kc'], 3)} (estimado)" if x["Kc_estimado"] else fmt(x["Kc"], 3) for x in res.salidas]),
        _fila("Cavitación incipiente xFzp1", "-",
              [f"{fmt(x['xFzp1'], 3)} (típico IEC)" if x["xFz_estimado"] else fmt(x["xFzp1"], 3)
               for x in res.salidas]),
    ]
    return filas


def _volumetrico(res: Resultado) -> bool:
    return all(c.Q is not None for c in res.dim.casos)


def _tabla_cierre(doc: "Documento", res: Resultado, pr: Presentacion):
    c = res.dim.cierre
    if c is None:
        return
    doc.h3(f"Condición de cierre ({c.nombre})")
    doc.tabla([["Presión de entrada P1", pr.vu(c.P1, "presion")],
               ["Presión de salida P2", pr.vu(c.P2, "presion")],
               ["ΔP de cierre", pr.vu(c.dP, "presion_dif")],
               ["Temperatura", pr.vu(c.T1, "temperatura") if c.T1 else "—"],
               ["Clase de cierre", res.inst.cierre]], anchos=[35, 65])


def _tabla_salidas(res: Resultado, pr: Presentacion, completa=True):
    s = res.salidas
    f = [
        _fila("Coeficiente de caudal Cv", "-", [x["Cv"] for x in s]),
        _fila("Apertura estimada", "%", [x["apertura"] * 100 for x in s], cifras=3),
        _fila("Carrera de operación", "", [x["carrera"] for x in s]),
        _fila("Coef. de cavitación de la aplicación Kapp", "-", [x["Kapp"] for x in s], cifras=3),
        _fila("Presión diferencial ΔP", "", [x["dP"] for x in s], pr, "presion_dif"),
        _fila("Relación ΔP/P1", "-", [x["dP_P1"] for x in s], cifras=3),
        _fila("ΔP de flujo estrangulado (ΔP choked)", "", [x["dPchoked"] for x in s], pr, "presion_dif"),
        _fila("ΔP de inicio de cavitación (ΔP cav)", "", [x["dPcav"] for x in s], pr, "presion_dif"),
        _fila("ΔP de cavitación incipiente (xFzp1)", "", [x["dPinc"] for x in s], pr, "presion_dif"),
        _fila("Condición", "", [x["condicion"] for x in s]),
    ]
    if completa:
        f += [
            _fila("Factor FF", "-", [x["FF"] for x in s], cifras=3),
            _fila("Factor de tubería Fp", "-", [x["FP"] for x in s], cifras=3),
            _fila("Factor combinado FLp", "-", [x["FLP"] for x in s], cifras=3),
            _fila("Viscosidad cinemática ν", "", [x["nu"] for x in s], pr, "viscosidad_cin"),
            _fila("Número de Reynolds Rev", "-", [x["Rev"] for x in s]),
            _fila("Densidad de entrada ρ1", "", [x["rho"] for x in s], pr, "densidad"),
            (_fila("Flujo másico", "", [x["W"] for x in s], pr, "caudal_masico") if _volumetrico(res) else
             _fila("Caudal volumétrico", "", [x["Q"] for x in s], pr, "caudal_vol")),
            _fila("Velocidad aguas arriba", "", [x["v1"] for x in s], pr, "velocidad", 3),
            _fila("Velocidad aguas abajo", "", [x["v2"] for x in s], pr, "velocidad", 3),
        ]
    f.append(_fila("Presión sonora a 1 m LpAe,1m", "dB(A)", [x["Lp"] for x in s], cifras=3))
    return f


def _tuberias(res: Resultado, pr: Presentacion):
    i = res.inst
    return [
        ["Entrada", tuberias.texto_nps(i.nps_entrada), i.sch_entrada, pr.vu(res.D1, "longitud"),
         pr.vu(res.tp1, "longitud", 3)],
        ["Salida", tuberias.texto_nps(i.nps_salida), i.sch_salida, pr.vu(res.D2, "longitud"),
         pr.vu(res.tp2, "longitud", 3)],
    ]


BALANCE = {True: "Balanceado", False: "Desbalanceado", None: "No aplica (cuerpo rotativo)"}


def _valvula_resumen(res: Resultado):
    i = res.inst
    return (f"{res.cuerpo.nombre}, {tuberias.texto_nps(res.nps_valvula)} (DN {res.informe.DN}), CL{i.rating}, "
            f"{i.conexion}; internos: {res.trim.nombre.lower()}; " + _texto_cv(res) + ".")


def _texto_cv(res: Resultado) -> str:
    tr = res.trim_reducido
    if tr is None:
        return f"Cv nominal típico {fmt(res.Cv_nominal)}"
    return (f"trim reducido, Cv {fmt(tr.Cv)} {'certificado' if tr.Cv_certificado else 'estimado por puerto'} "
            f"(trim completo típico: Cv {fmt(tr.Cv_tipico)})")


def _filas_cv(res: Resultado) -> list[list[str]]:
    """Cv nominal del trim completo típico y, si corresponde, del trim reducido."""
    tr = res.trim_reducido
    if tr is None:
        return [["Cv nominal (típico del tamaño y trim)", fmt(res.Cv_nominal)]]
    return [["Trim completo típico: puerto y Cv",
             f"{G.fraccion_in(tr.puerto_tipico_mm / 25.4)} ({G._mm(tr.puerto_tipico_mm)}) · Cv {fmt(tr.Cv_tipico)}"],
            ["Trim reducido: tamaño del puerto",
             f"{G.fraccion_in(tr.puerto_mm / 25.4)} ({G._mm(tr.puerto_mm)}) "
             + ("certificado" if tr.puerto_certificado else "estimado con el Cv certificado")],
            ["Cv nominal con el puerto reducido (usado en el cálculo)", tr.texto_Cv]]


# ---------------------------------------------------------------------------
# Memoria de cálculo
# ---------------------------------------------------------------------------
def memoria(res: Resultado, pr: Presentacion) -> Documento:
    if res.dim.compresible:
        return memoria_compresible(res, pr)
    doc = Documento("Memoria de cálculo", "memoria", res)
    d, s = res.dim, res.salidas
    casos = _casos(res)
    doc.h1("Memoria de cálculo de válvula de control")
    _identificacion(doc)
    doc.h2("1. Bases de cálculo")
    doc.lista([
        "Dimensionamiento de líquidos según IEC 60534-2-1:2011 / ANSI/ISA-75.01.01 (cláusulas 5, 7 y 9).",
        "Inicio de cavitación y ruido hidrodinámico según IEC 60534-8-4.",
        "Diámetros interiores de tubería según ASME B36.10M / B36.19M.",
        "Factores FL y Fd por caso: típicos de IEC 60534-2-1 (Tabla 2) o ingresados por el usuario. "
        "Kc: típico del trim o certificado; si se deja en blanco se estima con IEC 60534-8-4.",
        "Capacidad nominal (Cv a carrera total): típica del estilo.",
    ])
    doc.h2("2. Datos de entrada")
    doc.p(f"Fluido: {_nombre_fluido(res)}.")
    doc.tabla(_tabla_entradas(res, pr), ["Dato", "Unidad", *casos])
    _nota_salinidad(doc, res)
    _tabla_cierre(doc, res, pr)
    doc.tabla([["Presión crítica Pc", pr.vu(d.Pc, "presion_abs")],
               ["Válvula", _valvula_resumen(res)],
               ["Característica inherente", CARACTERISTICA.get(d.caracteristica, d.caracteristica)]],
              anchos=[35, 65])
    doc.tabla(_tuberias(res, pr), ["Tubería", "NPS", "Cédula", "Diámetro interior", "Espesor"])

    doc.h2("3. Cálculo")
    pd = pr.un("presion_dif")
    filas = [
        ["ΔP = P1 − P2", pd, *[pr.v(x["dP"], "presion_dif") for x in s]],
        ["FF = 0,96 − 0,28·√(Pv/Pc)", "-", *[fmt(x["FF"], 3) for x in s]],
        ["Fp = 1/√(1 + Σζ/N2·(C/d²)²)", "-", *[fmt(x["FP"], 3) for x in s]],
        ["FLp = FL/√(1 + FL²/N2·Σζ1·(C/d²)²)", "-", *[fmt(x["FLP"], 3) for x in s]],
        ["ΔPchoked = (FLp/Fp)²·(P1 − FF·Pv)", pd, *[pr.v(x["dPchoked"], "presion_dif") for x in s]],
        ["ΔP de dimensionamiento = mín(ΔP, ΔPchoked)", pd,
         *[pr.v(min(x["dP"], x["dPchoked"]), "presion_dif") for x in s]],
        ["Rev = N4·Fd·Q/(ν·√(C·FL))·(FL²·C²/(N2·D⁴) + 1)^¼", "-", *[fmt(x["Rev"]) for x in s]],
        ["FR (factor de Reynolds)", "-", *[fmt(x["FR"], 3) for x in s]],
        ["Cv = Q/(N1·Fp)·√(SG/ΔPdim)/0,865 (÷FR si Rev < 10⁴)", "-", *[fmt(x["Cv"]) for x in s]],
        ["Kapp = ΔP/(P1 − Pv)", "-", *[fmt(x["Kapp"], 3) for x in s]],
        ["Kc (cavitación constante del trim)", "-", *[fmt(x["Kc"], 3) for x in s]],
        ["ΔPcav = Kc·(P1 − Pv)", pd, *[pr.v(x["dPcav"], "presion_dif") for x in s]],
        ["xFzp1 = xFz·(6·10⁵/P1)^0,125 (cavitación incipiente)", "-", *[fmt(x["xFzp1"], 3) for x in s]],
        ["Apertura para Cv requerido", "%", *[fmt(x["apertura"] * 100, 3) for x in s]],
        [_ecuacion_carrera(res), "", *[x["carrera"] for x in s]],
        ["LpAe,1m (IEC 60534-8-4)", "dB(A)", *[fmt(x["Lp"], 3) for x in s]],
    ]
    doc.tabla(filas, ["Ecuación", "Unidad", *casos], resaltar_col=2)
    doc.nota("Constantes de IEC 60534-2-1, Tabla 1 (Kv en m³/h, ΔP en kPa, d en mm): N1 = 0,1; N2 = 1,6·10⁻³; "
             "N4 = 7,07·10⁻². Cv = Kv / 0,865.")
    doc.h2("4. Resultados")
    doc.tabla(_tabla_salidas(res, pr, completa=False), ["Resultado", "Unidad", *casos], resaltar_col=2)
    doc.grafico()
    doc.nota("Curva de Cv en función de la apertura de la válvula seleccionada (característica inherente) con "
             "los puntos de operación de cada caso.")
    _seccion_materiales_memoria(doc, res, 5)
    _seccion_equipo_memoria(doc, res, pr, 6)
    doc.h2("7. Conclusión")
    doc.p(f"Válvula seleccionada: {_valvula_resumen(res)}")
    criticas = [a for _, a in res.alertas if a.severidad >= Severidad.ADVERTENCIA]
    if criticas:
        titulos = sorted({a.titulo for a in criticas})
        doc.p("Condiciones a resolver: " + "; ".join(titulos) + ". Ver el informe técnico.")
    else:
        doc.p("No se detectaron condiciones que requieran acción.")
    return doc


# ---------------------------------------------------------------------------
# Hoja de datos
# ---------------------------------------------------------------------------
def hoja_datos(res: Resultado, pr: Presentacion) -> Documento:
    if res.dim.compresible:
        return hoja_datos_compresible(res, pr)
    doc = Documento("Hoja de datos", "hoja_datos", res)
    i, d, s = res.inst, res.dim, res.salidas
    casos = _casos(res)
    doc.h1("Hoja de datos de válvula de control")
    _identificacion(doc)
    doc.h2("Condiciones de servicio")
    doc.tabla([["Fluido", _nombre_fluido(res), "Estado", "Líquido"]], anchos=[18, 32, 18, 32])
    filas = _tabla_entradas(res, pr) + [
        _fila("Presión crítica Pc", "", [d.Pc] * len(s), pr, "presion_abs"),
        _fila("Presión diferencial ΔP", "", [x["dP"] for x in s], pr, "presion_dif"),
        _fila("Cv requerido", "-", [x["Cv"] for x in s]),
        _fila("Apertura", "%", [x["apertura"] * 100 for x in s], cifras=3),
        _fila("Carrera de operación", "", [x["carrera"] for x in s]),
        _fila("Kapp / Kc", "-", [f"{fmt(x['Kapp'], 3)} / {fmt(x['Kc'], 3)}" for x in s]),
        _fila("Condición", "", [x["condicion"] for x in s]),
        _fila("Ruido LpAe,1m", "dB(A)", [x["Lp"] for x in s], cifras=3),
        _fila("Velocidad aguas abajo", "", [x["v2"] for x in s], pr, "velocidad", 3),
    ]
    doc.tabla(filas, ["Dato", "Unidad", *casos])
    _nota_salinidad(doc, res)
    _tabla_cierre(doc, res, pr)
    doc.h2("Línea")
    doc.tabla(_tuberias(res, pr), ["Tubería", "NPS", "Cédula", "Diámetro interior", "Espesor"])
    doc.h2("Cuerpo")
    doc.tabla([
        ["Estilo", res.cuerpo.nombre],
        ["Diámetro", f"{tuberias.texto_nps(res.nps_valvula)} (DN {res.informe.DN})"],
        ["Rating", f"CL{i.rating} (ASME B16.34)"],
        ["Conexión a proceso", i.conexion],
        ["Presión de diseño", pr.vu(res.P_diseno, "presion")],
        ["Temperatura de diseño", pr.vu(res.T_diseno, "temperatura")],
        *_filas_material_cuerpo(res),
        ["Ruido admisible", f"{fmt(i.ruido_permitido, 3)} dB(A)"],
    ], anchos=[35, 65])
    doc.h2("Internos")
    rec = dict(res.recomendaciones)
    doc.tabla([
        ["Tipo de trim", res.trim.nombre],
        ["Obturador", BALANCE[d.balanceado]],
        ["Característica inherente", CARACTERISTICA.get(d.caracteristica, d.caracteristica)],
        *_filas_cv(res),
        *_filas_geometria(res),
        ["Guía y obturador", rec.get("Tipo de guía y obturador", "")],
        ["Clase de cierre", f"{i.cierre} ({REF_CIERRE})"],
        ["Fuga admisible", FUGA_CIERRE[i.cierre]],
        *_filas_material_internos(res),
    ], anchos=[35, 65])
    _seccion_equipo_hoja(doc, res, pr)
    doc.h2("Notas")
    notas = [f"[{NOMBRE_SEV[a.severidad]}] {caso}: {a.titulo}" for caso, a in res.alertas
             if a.severidad >= Severidad.ADVERTENCIA]
    doc.lista(notas or ["Sin condiciones que requieran acción."])
    doc.nota("Factores de válvula típicos y genéricos (no corresponden a ningún fabricante). Verificar con los "
             "datos certificados de la válvula ofertada.")
    return doc


# ---------------------------------------------------------------------------
# Informe técnico
# ---------------------------------------------------------------------------
def informe_tecnico(res: Resultado, pr: Presentacion) -> Documento:
    doc = Documento("Informe técnico", "informe_tecnico", res)
    casos = _casos(res)
    doc.h1("Informe técnico de selección de válvula de control")
    _identificacion(doc)
    peor = res.peor()
    doc.h2("1. Resumen")
    estado = {Severidad.INFO: "La selección no presenta condiciones que requieran acción.",
              Severidad.ADVERTENCIA: "La selección es viable con advertencias que deben revisarse.",
              Severidad.CRITICA: "Hay condiciones críticas: la selección requiere acciones de corrección."}[peor]
    doc.p(f"Válvula: {_valvula_resumen(res)}")
    doc.p(estado)
    doc.h2("2. Datos de instalación")
    i = res.inst
    doc.tabla([["Estilo de válvula", res.cuerpo.nombre], ["Rating", f"CL{i.rating}"],
               ["Conexión a proceso", i.conexion], ["Clase de cierre", i.cierre],
               ["Presión de diseño", pr.vu(res.P_diseno, "presion")],
               ["Temperatura de diseño", pr.vu(res.T_diseno, "temperatura")],
               ["Ruido permitido", f"{fmt(i.ruido_permitido, 3)} dB(A)"]], anchos=[35, 65])
    doc.tabla(_tuberias(res, pr), ["Tubería", "NPS", "Cédula", "Diámetro interior", "Espesor"])
    doc.h2("3. Datos de dimensionamiento")
    doc.p(f"Tipo de fluido: {CG.TIPOS[res.dim.tipo]}. Fluido: {_nombre_fluido(res)}.")
    doc.tabla(_entradas(res, pr), ["Dato", "Unidad", *casos])
    if res.dim.compresible:
        _datos_ruido(doc, res, pr)
    _tabla_cierre(doc, res, pr)
    doc.h2("4. Resultados")
    if res.dim.compresible:
        _salidas_compresible(doc, res, pr)
    else:
        doc.tabla(_tabla_salidas(res, pr), ["Resultado", "Unidad", *casos], resaltar_col=2)
    doc.grafico()
    doc.h2("5. Alertas y fenómenos")
    if not res.alertas:
        doc.p("Sin alertas.")
    for caso, a in res.alertas:
        doc.alerta(caso, a)
    doc.h2("6. Acciones de corrección")
    acciones, vistos = [], set()
    for _, a in res.alertas:
        if a.severidad < Severidad.ADVERTENCIA:
            continue
        for r in a.remedios:
            if r not in vistos:
                vistos.add(r)
                acciones.append(f"{a.titulo}: {r}")
    doc.lista(acciones or ["No se requieren acciones de corrección."])
    doc.h2("7. Características recomendadas de la válvula")
    doc.tabla([[t, x] for t, x in res.recomendaciones], ["Aspecto", "Recomendación"], anchos=[25, 75])
    _seccion_materiales_informe(doc, res, 8)
    _seccion_equipo_informe(doc, res, pr, 9)
    doc.h2("10. Recomendaciones finales")
    factores = "Cv, xT, FL y Fd" if res.dim.compresible else "Cv, FL, Fd y Kc"
    finales = [f"Solicitar al fabricante los valores certificados de {factores} del trim ofertado y "
               "repetir la verificación con esos datos."]
    if peor >= Severidad.CRITICA:
        finales.insert(0, "No emitir la compra hasta resolver las condiciones críticas indicadas en la sección 5.")
    if not res.dim.Cv_nominal or (res.trim_reducido and not res.trim_reducido.Cv_certificado):
        finales.append("La capacidad nominal usada es típica del estilo: confirmar la apertura con la curva real.")
    if res.materiales is not None and res.materiales.grupo_cuerpo:
        finales.append(f"Rating verificado con el material del cuerpo {res.materiales.cuerpo.nombre} (ASME B16.34, "
                       f"grupo {res.materiales.grupo_cuerpo}). Confirmar los materiales con el fabricante y con "
                       "el especialista en corrosión del proyecto.")
    else:
        finales.append("Verificar el rating con el material definitivo del cuerpo (la verificación actual es "
                       "referencial con acero al carbono, grupo 1.1).")
    if getattr(res, "equipo", None) is not None:
        finales.append("Confirmar con el fabricante del actuador el modelo, el resorte y el ajuste de banco: el "
                       "programa dimensiona con series genéricas y factores típicos de la industria.")
    doc.lista(finales)
    doc.h2("11. Alcance y referencias")
    doc.lista(["IEC 60534-2-1:2011 / ANSI/ISA-75.01.01: ecuaciones de dimensionamiento.",
               *_referencias_fluido(res),
               "ANSI/FCI 70-2 / IEC 60534-4: clases de cierre.",
               "ASME B16.34: clases de presión-temperatura. ASME B36.10M / B36.19M: tuberías.",
               "NACE MR0175 / ISO 15156 y NACE MR0103 / ISO 17945: materiales para servicio agrio.",
               "ISA-RP75.23: regímenes de cavitación. ASTM G98: resistencia al galling.",
               "Actuador: método de fuerzas A + B + C + D (vástago deslizante) y método de pares TB = A·ΔP + B, "
               "TD = C·ΔPef (rotativas), práctica de la industria de válvulas de control.",
               f"Calculado con {PRODUCTO} {__version__} de {EMPRESA}. Software independiente de fabricantes."])
    return doc


# ---------------------------------------------------------------------------
# Gas, vapor y vapor de agua
# ---------------------------------------------------------------------------
def _nombre_fluido(res: Resultado) -> str:
    d = res.dim
    if d.tipo == "steam":
        return "Agua, vapor (propiedades IAPWS-95)"
    t = d.fluido or "Definido por el usuario"
    if d.variante:
        t += f" — {d.variante}"
    if d.agua_libre and d.agua_libre != "Sí" and d.tipo != "liquido":
        t += f" (agua libre: {d.agua_libre.lower()})"
    return t


def _nota_salinidad(doc: "Documento", res: Resultado):
    from .fluido_servicio import correccion_salina
    c = correccion_salina(res.dim.salinidad_pct)
    if c:
        doc.nota(f"Propiedades del agua pura (IAPWS-95) corregidas por {c.texto} (correlaciones orientativas "
                 "para soluciones de NaCl).")


def _referencias_fluido(res: Resultado) -> list[str]:
    if not res.dim.compresible:
        return ["IEC 60534-8-4: ruido hidrodinámico e inicio de cavitación."]
    refs = ["IEC 60534-8-3:2010: ruido aerodinámico (trim, salida de la válvula y transmisión por la tubería).",
            "IEC 60534-2-1, Anexo A: Fd de trims multiorificio (Fd = 1/√No)."]
    if res.dim.tipo == "steam":
        refs.append("IAPWS-95 (formulación científica del agua y el vapor), mediante la biblioteca CoolProp.")
    elif res.dim.fluido:
        refs.append("Propiedades del fluido: ecuaciones de estado de referencia mediante la biblioteca CoolProp.")
    return refs


def _entradas(res: Resultado, pr: Presentacion):
    return _tabla_entradas_compresible(res, pr) if res.dim.compresible else _tabla_entradas(res, pr)


def _flujo_volumetrico_g(res: Resultado) -> bool:
    return res.dim.tipo != "steam" and all(c.QN is not None for c in res.dim.casos)


def _tabla_entradas_compresible(res: Resultado, pr: Presentacion):
    d, s, t = res.dim, res.salidas, res.dim.tipo
    filas = []
    if _flujo_volumetrico_g(res):
        filas.append(_fila("Caudal volumétrico estándar", "", [c.QN for c in d.casos], pr, "caudal_gas_std"))
    else:
        filas.append(_fila("Caudal másico", "", [x["W"] for x in s], pr, "caudal_masico"))
    filas += [_fila("Presión de entrada P1", "", [c.P1 for c in d.casos], pr, "presion"),
              _fila("Presión de salida P2", "", [c.P2 for c in d.casos], pr, "presion")]
    T1 = [pr.v(x["T1"], "temperatura") + (" (sat.)" if x.get("saturado") else "") for x in s]
    filas.append(["Temperatura de entrada T1", pr.un("temperatura"), *T1])
    if t in ("gas", "vapor"):
        filas.append(_fila("Masa molar / grav. específica", "", [x["M"] for x in s], pr, "masa_molar"))
    if t == "vapor":
        filas.append(_fila("Densidad de entrada ρ1", "", [x["rho"] for x in s], pr, "densidad"))
    if t in ("gas", "vapor"):
        filas += [_fila("Viscosidad dinámica μ", "", [x["mu"] for x in s], pr, "viscosidad_din"),
                  _fila("Relación de calores específicos γ", "-", [x["gamma"] for x in s])]
    if t == "gas":
        filas.append(_fila("Factor de compresibilidad Z1", "-", [x["Z1"] for x in s]))
    filas += [_fila("Factor de caída de presión crítica xT", "-", [c.xT for c in d.casos], cifras=3),
              _fila("Factor de recuperación FL", "-", [c.FL for c in d.casos], cifras=3),
              _fila("Modificador de estilo Fd", "-", [c.Fd for c in d.casos], cifras=3),
              _fila("Temperatura de salida T2", "", [x["T2"] for x in s], pr, "temperatura")]
    if t == "gas":
        filas.append(_fila("Factor de compresibilidad Z2", "-", [x["Z2"] for x in s]))
    if t == "vapor":
        filas.append(_fila("Densidad de salida ρ2", "", [x["rho2"] for x in s], pr, "densidad"))
    return filas


def _datos_ruido(doc: Documento, res: Resultado, pr: Presentacion):
    r, s = res.dim.ruido, res.salidas
    doc.tabla([["Trim para ruido aerodinámico", CG.TRIMS_RUIDO[r.trim][0]],
               ["Distancia Rn desde la pared de la tubería", pr.vu(r.Rn, "distancia", 3)],
               ["Área de salida de la válvula", pr.vu(s[0]["A_salida"], "area") +
                ("" if r.area_salida else " (área del diámetro nominal)")],
               ["Condición de salida T2, ρ2", s[0]["origen_salida"] or "dato del usuario"]], anchos=[40, 60])


def _salidas_compresible(doc: Documento, res: Resultado, pr: Presentacion, completa: bool = True):
    s, t = res.salidas, res.dim.tipo
    casos = _casos(res)
    enc = ["Resultado", "Unidad", *casos]
    if _flujo_volumetrico_g(res):
        otro = _fila("Flujo másico", "", [x["W"] for x in s], pr, "caudal_masico")
    elif t == "steam":
        otro = _fila("Caudal volumétrico real (entrada)", "", [x["Q1"] for x in s], pr, "caudal_vol")
    else:
        otro = _fila("Caudal volumétrico estándar", "", [x["QN"] for x in s], pr, "caudal_gas_std")
    dim = [_fila("Coeficiente de caudal Cv", "-", [x["Cv"] for x in s]),
           _fila("Apertura estimada", "%", [x["apertura"] * 100 for x in s], cifras=3),
        _fila("Carrera de operación", "", [x["carrera"] for x in s]),
           otro,
           _fila("Presión diferencial ΔP", "", [x["dP"] for x in s], pr, "presion_dif"),
           _fila("Relación de caída de presión x = ΔP/P1", "-", [x["x"] for x in s], cifras=3),
           _fila("ΔP de flujo estrangulado (ΔP choked)", "", [x["dPchoked"] for x in s], pr, "presion_dif"),
           _fila("Factor de tubería Fp", "-", [x["FP"] for x in s], cifras=3),
           _fila("Factor de recuperación combinado FLp", "-", [x["FLP"] for x in s], cifras=3),
           _fila("Factor xT ajustado xTP", "-", [x["xTP"] for x in s], cifras=3)]
    if completa:
        if t == "steam":
            dim += [_fila("Masa molar", "", [x["M"] for x in s], pr, "masa_molar"),
                    _fila("Relación de calores específicos γ", "-", [x["gamma"] for x in s])]
        if t in ("gas", "steam"):
            dim.append(_fila("Densidad de entrada ρ1", "", [x["rho"] for x in s], pr, "densidad"))
        if t in ("gas", "vapor"):
            dim.append(_fila("Viscosidad cinemática ν", "", [x["nu"] for x in s], pr, "viscosidad_cin"))
        dim.append(_fila("Factor de expansión Y", "-", [x["Y"] for x in s], cifras=3))
        if t == "steam":
            dim.append(_fila("Sobrecalentamiento a P2", "", [x["sobrecalentamiento2"] for x in s], pr,
                             "temperatura_dif", 3))
        dim += [_fila("Número de Reynolds Rev", "-", [x["Rev"] for x in s]),
                _fila("Diámetro interno aguas arriba", "", [x["D1"] for x in s], pr, "longitud"),
                _fila("Diámetro interno aguas abajo", "", [x["D2"] for x in s], pr, "longitud")]
    dim.append(_fila("Tipo de cálculo", "", [x["tipo_calculo"] for x in s]))
    doc.h3("Dimensionamiento")
    doc.tabla(dim, enc, resaltar_col=2)
    rui = [_fila("Nivel de trim de bajo ruido requerido", "", [x["nivel_trim"] for x in s]),
           _fila("Ruido del trim LpA a 1 m", "dB(A)", [x["Lp_trim"] for x in s], cifras=3),
           _fila("Ruido de la salida LpA a 1 m", "dB(A)", [x["Lp_salida"] for x in s], cifras=3),
           _fila("Ruido de la válvula LpA a 1 m", "dB(A)", [x["Lp"] for x in s], cifras=3),
           _fila(f"Ruido de la válvula LpA a Rn ({pr.vu(res.dim.ruido.Rn, 'distancia', 3)})", "dB(A)",
                 [x["Lp_Rn"] for x in s], cifras=3)]
    if t in ("gas", "steam"):
        rui.append(_fila("Densidad de salida ρ2", "", [x["rho2"] for x in s], pr, "densidad"))
    doc.h3("Ruido aerodinámico (IEC 60534-8-3)")
    doc.tabla(rui, enc, resaltar_col=2)
    if completa:
        vel = [_fila("Número de Mach aguas arriba", "-", [x["Ma1"] for x in s], cifras=3),
               _fila("Número de Mach en la salida de la válvula", "-", [x["Ma_valvula"] for x in s], cifras=3),
               _fila("Número de Mach aguas abajo", "-", [x["Ma2"] for x in s], cifras=3),
               _fila("Velocidad aguas arriba", "", [x["v1"] for x in s], pr, "velocidad", 3),
               _fila("Velocidad en la salida de la válvula", "", [x["v_valvula"] for x in s], pr, "velocidad", 3),
               _fila("Velocidad aguas abajo", "", [x["v2"] for x in s], pr, "velocidad", 3),
               _fila("Velocidad del sonido aguas arriba", "", [x["c1"] for x in s], pr, "velocidad"),
               _fila("Velocidad del sonido en la salida", "", [x["c_valvula"] for x in s], pr, "velocidad"),
               _fila("Velocidad del sonido aguas abajo", "", [x["c2"] for x in s], pr, "velocidad"),
               _fila("Área de salida de la válvula usada", "", [x["A_salida"] for x in s], pr, "area", 4)]
        doc.h3("Velocidades")
        doc.tabla(vel, enc, resaltar_col=2)


def memoria_compresible(res: Resultado, pr: Presentacion) -> Documento:
    doc = Documento("Memoria de cálculo", "memoria", res)
    d, s, t = res.dim, res.salidas, res.dim.tipo
    casos = _casos(res)
    doc.h1("Memoria de cálculo de válvula de control")
    _identificacion(doc)
    doc.h2("1. Bases de cálculo")
    bases = [f"Fluido compresible ({CG.TIPOS[t].lower()}): dimensionamiento según IEC 60534-2-1:2011 / "
             "ANSI/ISA-75.01.01 (cláusulas 6, 7, 8 y 9), con factor de expansión Y, factor de relación de "
             "calores específicos Fγ y corrección por accesorios (FP, xTP).",
             "Ruido aerodinámico según IEC 60534-8-3:2010 a 1 m aguas abajo y a 1 m de la pared de la tubería, "
             "separado en ruido del trim y de la salida de la válvula, y referido a la distancia Rn.",
             "Diámetros interiores y espesores de tubería según ASME B36.10M / B36.19M.",
             "Factores xT, FL y Fd por caso: típicos de IEC 60534-2-1 (Tabla 2) o ingresados por el usuario.",
             "Capacidad nominal (Cv a carrera total): típica del estilo."]
    if t == "steam":
        bases.insert(1, "Propiedades del vapor de agua según IAPWS-95; condición de salida por expansión "
                        "isoentálpica (estrangulamiento).")
    doc.lista(bases)
    doc.h2("2. Datos de entrada")
    doc.p(f"Fluido: {_nombre_fluido(res)}.")
    doc.tabla(_tabla_entradas_compresible(res, pr), ["Dato", "Unidad", *casos])
    _datos_ruido(doc, res, pr)
    _tabla_cierre(doc, res, pr)
    doc.tabla([["Válvula", _valvula_resumen(res)],
               ["Característica inherente", CARACTERISTICA.get(d.caracteristica, d.caracteristica)]],
              anchos=[35, 65])
    doc.tabla(_tuberias(res, pr), ["Tubería", "NPS", "Cédula", "Diámetro interior", "Espesor"])

    doc.h2("3. Cálculo")
    pd = pr.un("presion_dif")
    filas = [
        ["x = ΔP/P1", "-", *[fmt(x["x"], 4) for x in s]],
        ["Fγ = γ/1,40", "-", *[fmt(x["Fgamma"], 4) for x in s]],
        ["Fp = 1/√(1 + Σζ/N2·(C/d²)²)", "-", *[fmt(x["FP"], 4) for x in s]],
        ["xTP = (xT/Fp²)/(1 + xT·ζi/N5·(C/d²)²)", "-", *[fmt(x["xTP"], 4) for x in s]],
        ["xchoked = Fγ·xTP", "-", *[fmt(x["xchoked"], 4) for x in s]],
        ["ΔPchoked = xchoked·P1", pd, *[pr.v(x["dPchoked"], "presion_dif") for x in s]],
        ["Y = 1 − mín(x, xchoked)/(3·xchoked)", "-", *[fmt(x["Y"], 4) for x in s]],
        ["ρ1 (dato, P1·M/(Z1·R·T1) o IAPWS-95)", pr.un("densidad"), *[pr.v(x["rho"], "densidad") for x in s]],
        ["Kv = W/(N6·Fp·Y·√(mín(x, xchoked)·P1·ρ1))", "m³/h", *[fmt(x["Kv"]) for x in s]],
        ["Rev = N4·Fd·Q/(ν·√(C·FL))·(FL²·C²/(N2·D⁴) + 1)^¼", "-", *[fmt(x["Rev"]) for x in s]],
        ["FR (factor de Reynolds)", "-", *[fmt(x["FR"], 3) for x in s]],
        ["Cv = Kv/0,865 (÷FR si Rev < 10⁴)", "-", *[fmt(x["Cv"]) for x in s]],
        ["Apertura para Cv requerido", "%", *[fmt(x["apertura"] * 100, 3) for x in s]],
        [_ecuacion_carrera(res), "", *[x["carrera"] for x in s]],
        ["c = √(γ·P/ρ) aguas arriba / aguas abajo", pr.un("velocidad"),
         *[f"{pr.v(x['c1'], 'velocidad')} / {pr.v(x['c2'], 'velocidad')}" for x in s]],
        ["Ma = v/c en la salida de la válvula", "-", *[fmt(x["Ma_valvula"], 3) for x in s]],
        ["LpAe,1m trim / salida (IEC 60534-8-3)", "dB(A)",
         *[f"{fmt(x['Lp_trim'], 3)} / {fmt(x['Lp_salida'], 3)}" for x in s]],
        ["LpAe,1m de la válvula", "dB(A)", *[fmt(x["Lp"], 3) for x in s]],
        ["LpA a Rn = LpAe,1m − 10·lg((De + 2·Rn)/(De + 2))", "dB(A)", *[fmt(x["Lp_Rn"], 3) for x in s]],
    ]
    doc.tabla(filas, ["Ecuación", "Unidad", *casos], resaltar_col=2)
    doc.nota("Constantes de IEC 60534-2-1, Tabla 1 (Kv en m³/h, P en kPa, ρ en kg/m³, W en kg/h, d en mm): "
             "N2 = 1,6·10⁻³; N4 = 7,07·10⁻²; N5 = 1,8·10⁻³; N6 = 3,16. De = Di + 2·tp (diámetro exterior de la "
             "tubería de salida). Cv = Kv / 0,865.")
    doc.h2("4. Resultados")
    _salidas_compresible(doc, res, pr, completa=False)
    doc.grafico()
    doc.nota("Curva de Cv en función de la apertura de la válvula seleccionada (característica inherente) con "
             "los puntos de operación de cada caso.")
    _seccion_materiales_memoria(doc, res, 5)
    _seccion_equipo_memoria(doc, res, pr, 6)
    doc.h2("7. Conclusión")
    doc.p(f"Válvula seleccionada: {_valvula_resumen(res)}")
    criticas = [a for _, a in res.alertas if a.severidad >= Severidad.ADVERTENCIA]
    if criticas:
        doc.p("Condiciones a resolver: " + "; ".join(sorted({a.titulo for a in criticas})) +
              ". Ver el informe técnico.")
    else:
        doc.p("No se detectaron condiciones que requieran acción.")
    return doc


def hoja_datos_compresible(res: Resultado, pr: Presentacion) -> Documento:
    doc = Documento("Hoja de datos", "hoja_datos", res)
    i, d, s = res.inst, res.dim, res.salidas
    casos = _casos(res)
    doc.h1("Hoja de datos de válvula de control")
    _identificacion(doc)
    doc.h2("Condiciones de servicio")
    doc.tabla([["Fluido", _nombre_fluido(res), "Estado", CG.TIPOS[d.tipo]]], anchos=[18, 32, 18, 32])
    filas = _tabla_entradas_compresible(res, pr) + [
        _fila("Presión diferencial ΔP", "", [x["dP"] for x in s], pr, "presion_dif"),
        _fila("Cv requerido", "-", [x["Cv"] for x in s]),
        _fila("Apertura", "%", [x["apertura"] * 100 for x in s], cifras=3),
        _fila("Carrera de operación", "", [x["carrera"] for x in s]),
        _fila("Tipo de cálculo", "", [x["tipo_calculo"] for x in s]),
        _fila("Ruido de la válvula LpA a 1 m", "dB(A)", [x["Lp"] for x in s], cifras=3),
        _fila("Nivel de trim de bajo ruido", "", [x["nivel_trim"] for x in s]),
        _fila("Mach en la salida de la válvula", "-", [x["Ma_valvula"] for x in s], cifras=3),
        _fila("Velocidad aguas abajo", "", [x["v2"] for x in s], pr, "velocidad", 3),
    ]
    doc.tabla(filas, ["Dato", "Unidad", *casos])
    _tabla_cierre(doc, res, pr)
    doc.h2("Línea")
    doc.tabla(_tuberias(res, pr), ["Tubería", "NPS", "Cédula", "Diámetro interior", "Espesor"])
    doc.h2("Cuerpo")
    doc.tabla([
        ["Estilo", res.cuerpo.nombre],
        ["Diámetro", f"{tuberias.texto_nps(res.nps_valvula)} (DN {res.informe.DN})"],
        ["Rating", f"CL{i.rating} (ASME B16.34)"],
        ["Conexión a proceso", i.conexion],
        ["Área de salida", pr.vu(s[0]["A_salida"], "area")],
        ["Presión de diseño", pr.vu(res.P_diseno, "presion")],
        ["Temperatura de diseño", pr.vu(res.T_diseno, "temperatura")],
        *_filas_material_cuerpo(res),
        ["Ruido admisible", f"{fmt(i.ruido_permitido, 3)} dB(A)"],
    ], anchos=[35, 65])
    doc.h2("Internos")
    rec = dict(res.recomendaciones)
    doc.tabla([
        ["Tipo de trim", res.trim.nombre],
        ["Obturador", BALANCE[d.balanceado]],
        ["Trim para ruido aerodinámico", CG.TRIMS_RUIDO[d.ruido.trim][0]],
        ["Característica inherente", CARACTERISTICA.get(d.caracteristica, d.caracteristica)],
        *_filas_cv(res),
        *_filas_geometria(res),
        ["Guía y obturador", rec.get("Tipo de guía y obturador", "")],
        ["Clase de cierre", f"{i.cierre} ({REF_CIERRE})"],
        ["Fuga admisible", FUGA_CIERRE[i.cierre]],
        *_filas_material_internos(res),
    ], anchos=[35, 65])
    _seccion_equipo_hoja(doc, res, pr)
    doc.h2("Notas")
    notas = [f"[{NOMBRE_SEV[a.severidad]}] {caso}: {a.titulo}" for caso, a in res.alertas
             if a.severidad >= Severidad.ADVERTENCIA]
    doc.lista(notas or ["Sin condiciones que requieran acción."])
    doc.nota("Factores de válvula típicos y genéricos (no corresponden a ningún fabricante). Verificar con los "
             "datos certificados de la válvula ofertada.")
    return doc


# ---------------------------------------------------------------------------
# Materiales (solapa 3)
# ---------------------------------------------------------------------------
_SIN_MAT = "A definir: complete la solapa 3-Materiales y recalcule"


def _comp_mat(res: Resultado, *claves) -> list[list[str]]:
    """Filas «componente / material (dureza)» para la hoja de datos."""
    m = res.materiales
    filas = []
    for k in claves:
        if m is None:
            continue
        rc = m.componentes.get(k)
        if rc is None or not rc.aplica:
            continue
        t = rc.texto
        if rc.elegido:
            t += f" · {rc.elegido.dureza_txt}"
            if rc.elegido.astm and rc.elegido.astm not in t:
                t += f" · {rc.elegido.astm}"
        if rc.modo == "MANUAL":
            t += " · selección MANUAL"
        filas.append([f"Material {rc.nombre.lower()}", t])
    return filas


def _filas_material_cuerpo(res: Resultado) -> list[list[str]]:
    if res.materiales is None:
        return [["Material del cuerpo", _SIN_MAT]]
    filas = _comp_mat(res, "cuerpo", "bonete")
    g = res.materiales.grupo_cuerpo
    if g:
        filas.append(["Grupo ASME B16.34", g])
    return filas


def _ecuacion_carrera(res: Resultado) -> str:
    g = res.geometria
    if g.rotativa:
        return f"Rotación = apertura × {g.rotacion_grados:.0f}° (rotación nominal)"
    return f"Carrera = apertura × carrera nominal ({g.texto_carrera}; puerto {g.texto_puerto})"


def _filas_geometria(res: Resultado) -> list[list[str]]:
    g = res.geometria
    nombres = [rc.caso.nombre for rc in res.informe.casos]
    op = " · ".join(f"{n}: {s['carrera']}" for n, s in zip(nombres, res.salidas))
    return [[g.nombre_puerto, g.texto_puerto],
            ["Rotación nominal del eje" if g.rotativa else "Carrera nominal del vástago", g.texto_carrera],
            ["Rotación de operación" if g.rotativa else "Carrera de operación", op]]


def _filas_material_internos(res: Resultado) -> list[list[str]]:
    m = res.materiales
    if m is None:
        return [["Material de internos", _SIN_MAT]]
    filas = _comp_mat(res, "obturador", "asiento", "jaula", "guia", "vastago")
    if m.juegos:
        filas.append(["Juego de internos", m.juego_texto])
    if m.agrio:
        filas.append(["Servicio agrio", f"Sí: materiales según {m.norma}"])
    e = m.empaquetadura
    if e is not None:
        filas += [["Empaquetadura", e.texto + (" · selección MANUAL" if e.modo == "MANUAL" else "")]]
        if e.elegida is not None:
            filas.append(["Tipo de empaquetadura", f"{e.elegida.tipo} · carga viva: "
                          f"{'sí' if e.elegida.carga_viva else 'no'} · {e.elegida.guia_temperatura}"])
        filas += [["Emisión fugitiva", f"Requerida {e.emision_req}" + (f"; {e.elegida.emision}" if e.elegida else "")],
                  ["Tipo de bonete", e.bonete]]
    return filas


def _seccion_materiales_informe(doc: Documento, res: Resultado, n: int):
    m = res.materiales
    doc.h2(f"{n}. Selección de materiales")
    if m is None:
        doc.p("No se definió el servicio de materiales (solapa 3-Materiales). Los materiales quedan a definir.")
        return
    doc.p(f"Medio detectado: {m.medio_detectado}. Servicio agrio: {'sí' if m.agrio else 'no'}"
          + (f" (pH₂S = {fmt(m.pH2S_kPa, 3)} kPa)" if m.pH2S_kPa is not None else "")
          + f". Norma aplicable: {m.norma} ({m.norma_motivo.rstrip('.')}).")
    doc.p(f"Dureza de internos: {m.dureza_motivo}")
    filas = []
    for rc in m.componentes.values():
        just = "; ".join(x.strip(" ·") for x in rc.justificacion[:3]) if rc.aplica else ""
        filas.append([rc.nombre, rc.modo if rc.aplica else "—", rc.texto,
                      rc.elegido.dureza_txt if (rc.aplica and rc.elegido) else "—", just])
    doc.tabla(filas, ["Componente", "Modo", "Material", "Dureza", "Justificación"], anchos=[13, 11, 26, 11, 39])
    if any(rc.modo == "MANUAL" for rc in m.componentes.values() if rc.aplica):
        doc.nota("Hay componentes en selección MANUAL: no recomendada. La responsabilidad por la aptitud del "
                 "material queda fuera del alcance del programa.")
    if m.galling:
        doc.h3("Riesgo de galling en piezas en contacto")
        doc.tabla([[g.par, f"{g.a} / {g.b}", g.nivel, g.motivo + (f" {g.remedio}" if g.nivel != "BAJO" else "")]
                   for g in m.galling], ["Pareja", "Materiales", "Riesgo", "Motivo y remedio"], anchos=[18, 20, 10, 52])
    alertas = m.todas_las_alertas()
    if alertas:
        doc.h3("Alertas de materiales y su resolución")
        doc.lista([f"[{NOMBRE_SEV[a.severidad]}] {a.titulo}" + (f": {a.remedios[0]}" if a.remedios else "")
                   for a in alertas])
    if m.juegos:
        doc.h3("Juego de internos")
        doc.p(m.juego_texto + ".")
        if m.juego is not None and m.juego.juego.aplicacion:
            doc.p("Aplicación típica: " + m.juego.juego.aplicacion)
        aptos = [v.juego.texto for v in sorted(m.juegos, key=lambda v: v.puntaje) if v.ok]
        doc.p("Juegos estándar aptos para este servicio: " + ("; ".join(aptos) if aptos else "ninguno") + ".")
    e = m.empaquetadura
    if e is not None:
        doc.h3("Empaquetadura y bonete")
        tipo = (f"{e.elegida.tipo}; carga viva: {'sí' if e.elegida.carga_viva else 'no'}; "
                f"{e.elegida.guia_temperatura}") if e.elegida else "—"
        doc.tabla([["Empaquetadura", f"{e.texto} ({e.modo})"], ["Tipo de empaquetadura", tipo],
                   ["Tipo de bonete", e.bonete],
                   ["Emisión requerida", f"{e.emision_req}: {e.emision_motivo}"],
                   ["Motivo del bonete", e.bonete_motivo]], anchos=[25, 75])
        if e.notas:
            doc.lista(e.notas)
    if m.notas:
        doc.lista(m.notas)
    doc.nota("Preselección de ingeniería: no reemplaza la aprobación del fabricante ni del especialista en "
             "corrosión del proyecto.")


def _dato_agente(a) -> str:
    e = a.entrada
    if e.concentracion is None:
        return "Sin concentración"
    return f"{fmt(e.concentracion, 4)} {e.unidad} ({e.fase})"


def _seccion_materiales_memoria(doc: Documento, res: Resultado, n: int):
    m = res.materiales
    doc.h2(f"{n}. Selección de materiales")
    if m is None:
        doc.p("Sin datos de servicio de materiales: los materiales quedan a definir.")
        return
    sv, ctx = m.servicio, m.ctx
    doc.h3(f"{n}.1 Bases")
    doc.lista([
        "Límites presión-temperatura: ASME B16.34-2017 (tablas 2, por grupo de material).",
        "Servicio agrio: NACE MR0175 / ISO 15156 (producción) y NACE MR0103 / ISO 17945 (refinación); umbral "
        "pH₂S ≥ 0,3 kPa con agua libre.",
        "Resistencia química: reglas por familia de material y extracto de tabla de compatibilidad A–D, "
        "combinadas tomando el nivel más desfavorable.",
        "Cavitación: regímenes ISA-RP75.23 con Kapp, Kc, FL² y xFzp1; dureza mínima de internos por régimen.",
        "Galling: criterios de pares de materiales (referencia de ensayo ASTM G98).",
        "Puntaje de selección: índice de costo × factor químico × factor de erosión; se elige el menor apto.",
    ])
    doc.h3(f"{n}.2 Servicio")
    doc.tabla([["Industria", sv.industria],
               ["Fluido (solapa 2)", f"{_nombre_fluido(res)} → familia {sv.familia_fluido.lower()}"],
               ["Fase (deducida del cálculo)", sv.fase], ["Composición", sv.composicion], ["Agua libre", sv.acuosa],
               ["Régimen", sv.regimen], ["Medio (usuario / detectado)", f"{sv.medio} / {m.medio_detectado}"],
               ["Sólidos", sv.solidos], ["Norma de servicio agrio", f"{sv.norma_agrio} → {m.norma}"],
               ["Temperatura de operación", f"{fmt(ctx.T_min_C, 3)} a {fmt(ctx.T_max_C, 3)} °C"],
               ["Presión máxima", f"{fmt(ctx.P_kPa, 4)} kPa abs"],
               ["ΔP máximo", f"{fmt(ctx.dP_max_kPa, 4)} kPa"],
               ["Cavitación / flashing", f"{ctx.severidad_cav} / {'sí' if ctx.flashing else 'no'}"]],
              anchos=[35, 65])
    if m.agentes:
        doc.h3(f"{n}.3 Agentes químicos normalizados")
        doc.tabla([[f"{a.agente.formula} — {a.agente.nombre}", _dato_agente(a), a.texto(),
                    ("Supuesto conservador. " if a.supuesto else "") + " ".join(a.notas)] for a in m.agentes],
                  ["Agente", "Dato ingresado", "Valor normalizado", "Nota"], anchos=[25, 20, 20, 35])
    if ctx.casos_cav:
        doc.h3(f"{n}.4 Cavitación por caso")
        doc.tabla([[c, sev, fmt(k, 3), fmt(kc, 3), fmt(fl2, 3)] for c, sev, k, kc, fl2 in ctx.casos_cav],
                  ["Caso", "Régimen", "Kapp", "Kc", "FL²"])
    doc.p(f"Dureza requerida: {m.dureza_motivo}")
    if m.dureza_req:
        doc.tabla([[k, f"{fmt(v, 3)} HRC"] for k, v in m.dureza_req.items()], ["Componente", "Dureza mínima"])
    if m.mecanismos:
        doc.h3(f"{n}.5 Mecanismos de deterioro")
        doc.tabla([[x.codigo, x.nombre, x.estado, x.motivo] for x in m.mecanismos],
                  ["ID", "Mecanismo", "Estado", "Motivo"], anchos=[7, 28, 12, 53])
    doc.h3(f"{n}.6 Selección por componente")
    for rc in m.componentes.values():
        if not rc.aplica:
            doc.p(f"{rc.nombre}: no aplica con el trim elegido.")
            continue
        doc.p(f"{rc.nombre} ({rc.modo}): {rc.texto}"
              + (f". Recomendado: {rc.recomendado.nombre}" if rc.recomendado and rc.recomendado is not rc.elegido
                 else "") + ".")
        items = [x.strip(" ·") for x in rc.justificacion]
        ev = rc.eval_elegido
        if ev is not None:
            items.append(f"Nivel químico {ev.nivel_quimico}; puntaje {fmt(ev.puntaje, 3)}"
                         + (f"; rating {fmt(ev.rating, 3)} bar g, clase mínima CL{ev.clase_min}"
                            if ev.rating is not None and ev.clase_min else ""))
            items += [f"{c.agente}: {c.nivel} — {c.detalle}" for c in ev.quimica]
        if rc.alternativas:
            items.append("Alternativas aptas: " + "; ".join(f"{a.nombre} ({b})" for a, b in rc.alternativas[:4]))
        if rc.descartados:
            items.append(f"Materiales descartados: {len({a.clave for a, _ in rc.descartados})} (detalle en la "
                         "solapa 3-Materiales, botón Análisis).")
        doc.lista(items)
    if m.galling:
        doc.h3(f"{n}.7 Galling")
        doc.tabla([[g.par, f"{g.a} / {g.b}", g.nivel, g.motivo, g.remedio] for g in m.galling],
                  ["Pareja", "Materiales", "Riesgo", "Motivo", "Remedio"], anchos=[15, 17, 9, 30, 29])
    if m.juegos:
        doc.h3(f"{n}.8 Juego de internos")
        doc.p("Cada juego estándar se evalúa con las mismas reglas que los materiales sueltos (temperatura, "
              "química, servicio agrio, dureza y resistencia al desgaste) más el galling de las parejas "
              "obturador–asiento y obturador–jaula y el tipo de asiento que pide la clase de cierre.")
        doc.p(m.juego_texto + ".")
        doc.lista([f"{v.juego.texto}: " + ("apto" if v.ok else "no apto. " + " ".join(v.motivos[:2]))
                   + f" (índice {fmt(v.puntaje, 3)})" for v in m.juegos])
    e = m.empaquetadura
    if e is not None:
        doc.h3(f"{n}.9 Empaquetadura y bonete")
        doc.p("Bases: límites de temperatura (con reducción del grafito en medio oxidante), presión, compatibilidad "
              "química A–D, vacío, vapor, sólidos y exigencia de emisión fugitiva (ISO 15848-1; API 622 / 624 "
              "para vástago ascendente, API 641 para cuarto de vuelta). Entre las aptas se elige el menor índice "
              "de costo × factor químico × factores de servicio.")
        doc.p(f"Elegida ({e.modo}): {e.texto}. Bonete: {e.bonete}. {e.bonete_motivo}")
        items = list(e.justificacion) + list(e.notas)
        if e.alternativas:
            items.append("Alternativas aptas: " + "; ".join(f"{a.nombre} ({b})" for a, b in e.alternativas[:4]))
        items += [f"Descartada: {a.nombre}. {b}" for a, b in e.descartadas]
        doc.lista(items)
    for t in m.notas:
        doc.nota(t)


CONSTRUCTORES = {"memoria": memoria, "hoja_datos": hoja_datos, "informe_tecnico": informe_tecnico}


# ---------------------------------------------------------------------------
# Salida HTML (vista previa e impresión)
# ---------------------------------------------------------------------------
CSS = """
body{font-family:'Segoe UI',Arial,sans-serif;font-size:10pt;color:#1d2733}
h1{font-size:15pt;color:#0d3b66;margin:6px 0}
h2{font-size:12pt;color:#0d3b66;border-bottom:1px solid #0d3b66;margin-top:14px}
h3{font-size:10.5pt;margin:8px 0 2px 0}
table{border-collapse:collapse;margin:4px 0 8px 0}
td,th{border:1px solid #9aa5b1;padding:3px 6px;vertical-align:top}
th{background:#0d3b66;color:white;font-weight:600}
td.r{background:#e3f1e6}
td.et{background:#eef2f6;font-weight:600}
p.nota{color:#55606d;font-size:8.5pt;font-style:italic}
.cab{background:#0d3b66;color:white;padding:6px}
"""


def a_html(doc: Documento, src_grafico: str = "grafico://cv", encabezado: bool = True) -> str:
    e = _html.escape
    h = []
    if encabezado:
        h.append(f"<table width='100%' class='cab' style='border:none'><tr>"
                 f"<td style='border:none;background:#0d3b66;color:white'><b>{EMPRESA}</b> · {PRODUCTO}</td>"
                 f"<td style='border:none;background:#0d3b66;color:white' align='right'>{e(doc.titulo)}"
                 f" · Tag {e(doc.res.ident.tag or '—')} · Rev. {e(doc.res.ident.revision or '—')}</td></tr></table>")
    if MARCA_AGUA:
        h.append(f"<p align='center' style='color:#c0392b;font-weight:bold;border:1px solid #c0392b;"
                 f"padding:4px'>{e(MARCA_AGUA)}</p>")
    for b in doc.bloques:
        t = b.tipo
        if t in ("h1", "h2", "h3"):
            h.append(f"<{t}>{e(b.contenido)}</{t}>")
        elif t == "p":
            h.append(f"<p>{e(b.contenido)}</p>")
        elif t == "nota":
            h.append(f"<p class='nota'>{e(b.contenido)}</p>")
        elif t == "lista":
            h.append("<ul>" + "".join(f"<li>{e(x)}</li>" for x in b.contenido) + "</ul>")
        elif t == "grafico":
            h.append(f"<p align='center'><img src='{src_grafico}' width='620'></p>")
        elif t == "tabla":
            enc, rc = b.opciones.get("encabezado"), b.opciones.get("resaltar_col")
            filas = []
            if enc:
                filas.append("<tr>" + "".join(f"<th>{e(c)}</th>" for c in enc) + "</tr>")
            for fila in b.contenido:
                celdas = []
                for j, c in enumerate(fila):
                    if enc is None and j % 2 == 0:
                        cls = " class='et'"
                    elif rc is not None and j >= rc:
                        cls = " class='r' align='right'"
                    elif enc and j >= 2:
                        cls = " align='right'"
                    else:
                        cls = ""
                    celdas.append(f"<td{cls}>{e(str(c))}</td>")
                filas.append("<tr>" + "".join(celdas) + "</tr>")
            h.append("<table width='100%' cellspacing='0' cellpadding='3'>" + "".join(filas) + "</table>")
        elif t == "alerta":
            caso, a = b.contenido
            col = COLOR_SEV[a.severidad]
            inten = ""
            if a.intensidad is not None:
                inten = f" · intensidad {a.intensidad:.0f} %" + (f" ({a.nivel})" if a.nivel else "")
            h.append(f"<h3 style='color:#{col}'>[{NOMBRE_SEV[a.severidad]}] {e(caso)}: {e(a.titulo)}{inten}</h3>"
                     f"<p>{e(a.mensaje)}</p>")
            if a.remedios:
                h.append("<p><b>Soluciones posibles:</b></p><ul>" + "".join(f"<li>{e(r)}</li>" for r in a.remedios)
                         + "</ul>")
            if a.referencia:
                h.append(f"<p class='nota'>Referencia: {e(a.referencia)}</p>")
        elif t == "salto":
            h.append("<hr>")
    return f"<html><head><meta charset='utf-8'><style>{CSS}</style></head><body>{''.join(h)}</body></html>"


# ---------------------------------------------------------------------------
# Salida Word (.docx editable)
# ---------------------------------------------------------------------------
def a_docx(docs: list[Documento], ruta: str, png_grafico: bytes | None = None):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    AZUL = RGBColor(0x0D, 0x3B, 0x66)

    def sombrear(celda, color):
        tcPr = celda._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), color)
        tcPr.append(shd)

    def campo(parrafo, instr):
        r = parrafo.add_run()
        for tipo, texto in (("begin", None), (None, instr), ("end", None)):
            if tipo:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tipo)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = texto
            r._r.append(el)

    w = Document()
    st = w.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    for nombre, tam in (("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 10.5)):
        hs = w.styles[nombre]
        hs.font.color.rgb = AZUL
        hs.font.size = Pt(tam)
        hs.font.name = "Calibri"
    sec = w.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.orientation = WD_ORIENT.PORTRAIT
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Cm(1.8))
    sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.6)
    ident = docs[0].res.ident
    cab = sec.header.paragraphs[0]
    cab.text = ""
    r = cab.add_run(f"{EMPRESA}  ·  {PRODUCTO}")
    r.bold = True
    r.font.color.rgb = AZUL
    cab.add_run(f"\t\tTag {ident.tag or '—'}  ·  Rev. {ident.revision or '—'}").font.size = Pt(8.5)
    pie = sec.footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pie.add_run(f"{PRODUCTO} {__version__} · Página ").font.size = Pt(8)
    campo(pie, "PAGE")
    pie.add_run(" de ").font.size = Pt(8)
    campo(pie, "NUMPAGES")
    if MARCA_AGUA:
        ROJO = RGBColor(0xC0, 0x39, 0x2B)
        for par in (sec.header.add_paragraph(), sec.footer.add_paragraph()):
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = par.add_run(MARCA_AGUA)
            r.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = ROJO
    ancho_util = 21.0 - 3.6

    for k, doc in enumerate(docs):
        if k:
            w.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        for b in doc.bloques:
            t = b.tipo
            if t == "h1":
                w.add_heading(b.contenido, 1)
            elif t == "h2":
                w.add_heading(b.contenido, 2)
            elif t == "h3":
                w.add_heading(b.contenido, 3)
            elif t == "p":
                w.add_paragraph(b.contenido)
            elif t == "nota":
                p = w.add_paragraph()
                r = p.add_run(b.contenido)
                r.italic = True
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(0x55, 0x60, 0x6D)
            elif t == "lista":
                for x in b.contenido:
                    w.add_paragraph(x, style="List Bullet")
            elif t == "grafico":
                if png_grafico:
                    p = w.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.add_run().add_picture(io.BytesIO(png_grafico), width=Cm(16))
            elif t == "tabla":
                enc, rc, anchos = (b.opciones.get(x) for x in ("encabezado", "resaltar_col", "anchos"))
                filas = ([enc] if enc else []) + b.contenido
                ncol = max(len(f) for f in filas)
                tb = w.add_table(rows=len(filas), cols=ncol)
                tb.style = "Table Grid"
                tb.alignment = WD_TABLE_ALIGNMENT.CENTER
                for i, fila in enumerate(filas):
                    for j in range(ncol):
                        c = tb.cell(i, j)
                        c.text = str(fila[j]) if j < len(fila) else ""
                        par = c.paragraphs[0]
                        for run in par.runs:
                            run.font.size = Pt(9)
                        if enc and i == 0:
                            sombrear(c, "0D3B66")
                            for run in par.runs:
                                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                                run.bold = True
                        elif enc is None and j % 2 == 0:
                            sombrear(c, "EEF2F6")
                            for run in par.runs:
                                run.bold = True
                        elif rc is not None and j >= rc:
                            sombrear(c, "E3F1E6")
                            par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                        elif enc and j >= 2:
                            par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if not anchos and enc and ncol >= 3 and enc[1] == "Unidad":
                    anchos = [38, 10] + [52 / (ncol - 2)] * (ncol - 2)
                if anchos:
                    tb.autofit = False
                    tot = sum(anchos)
                    for j, a in enumerate(anchos):
                        for fila in tb.rows:
                            fila.cells[j].width = Cm(ancho_util * a / tot)
                w.add_paragraph()
            elif t == "alerta":
                caso, a = b.contenido
                p = w.add_paragraph()
                inten = ""
                if a.intensidad is not None:
                    inten = f" · intensidad {a.intensidad:.0f} %" + (f" ({a.nivel})" if a.nivel else "")
                r = p.add_run(f"[{NOMBRE_SEV[a.severidad]}] {caso}: {a.titulo}{inten}")
                r.bold = True
                r.font.color.rgb = RGBColor.from_string(COLOR_SEV[a.severidad])
                w.add_paragraph(a.mensaje)
                if a.remedios:
                    w.add_paragraph().add_run("Soluciones posibles:").bold = True
                    for x in a.remedios:
                        w.add_paragraph(x, style="List Bullet")
                if a.referencia:
                    r = w.add_paragraph().add_run(f"Referencia: {a.referencia}")
                    r.italic = True
                    r.font.size = Pt(8.5)
    w.core_properties.title = " / ".join(d.titulo for d in docs)
    w.core_properties.author = ident.preparado or EMPRESA
    w.save(ruta)


# ---------------------------------------------------------------------------
# Actuador, posicionador y accesorios (solapas 4, 5 y 6)
# ---------------------------------------------------------------------------
_SIN_EQUIPO = "A definir: complete las solapas 4-Actuador, 5-Posicionador y 6-Accesorios y recalcule"


def _aire_texto(pr: Presentacion, act) -> str:
    if act.P_aire_min is None:
        return "No informado"
    u = "bar g" if pr.un("presion") in ("bar a", "bar g", "kPa a", "kPa g", "MPa a", "MPa g") else "psig"
    a = U.desde_base(act.P_aire_min + 101.325, "presion", u)
    t = f"{fmt(a, 3)}"
    if act.P_aire_max is not None:
        t += f" a {fmt(U.desde_base(act.P_aire_max + 101.325, 'presion', u), 3)}"
    return f"{t} {u} (mínimo a máximo)"


def _filas_actuador(res: Resultado, pr: Presentacion) -> list[list[str]]:
    eq = res.equipo
    act, ra = eq.datos.actuador, eq.actuador
    filas = [["Tipo de actuador", act.tipo], ["Modo de falla", act.falla],
             ["Volante (mando manual)", "Sí" if act.volante else "No"],
             ["Aire de instrumentos disponible", _aire_texto(pr, act)]]
    if not ra.rotativa:
        filas.append(["Flujo (sentido)", act.flujo])
    filas.append(["Selección", ra.seleccion or "—"])
    filas += [[f, v] for f, v, _ in ra.filas if f not in ("Tamaño (serie genérica)",)]
    return filas


def _seccion_equipo_hoja(doc: Documento, res: Resultado, pr: Presentacion):
    eq = getattr(res, "equipo", None)
    doc.h2("Actuador")
    if eq is None:
        doc.tabla([["Actuador", _SIN_EQUIPO], ["Posicionador y accesorios", _SIN_EQUIPO]], anchos=[35, 65])
        return
    ra = eq.actuador
    filas = _filas_actuador(res, pr)
    if ra.requerido:
        from . import actuadores as A
        filas.insert(5, ["Par requerido (con margen 1,25)" if ra.rotativa else "Fuerza requerida (con margen 1,10)",
                         A._Nm(ra.requerido) if ra.rotativa else A._N(ra.requerido)])
    doc.tabla(filas, anchos=[35, 65])
    doc.h2("Posicionador")
    doc.tabla([[f, v] for f, v, _ in eq.posicionador.filas], anchos=[35, 65])
    doc.h2("Accesorios")
    doc.tabla([[f, v] for f, v, _ in eq.accesorios.filas] or [["Accesorios", "No requeridos"]], anchos=[35, 65])


def _seccion_equipo_memoria(doc: Documento, res: Resultado, pr: Presentacion, n: int):
    eq = getattr(res, "equipo", None)
    doc.h2(f"{n}. Actuador, posicionador y accesorios")
    if eq is None:
        doc.p(_SIN_EQUIPO + ".")
        return
    ra = eq.actuador
    if ra.rotativa:
        doc.p("Método de pares para válvulas rotativas: par de arranque TB = A·ΔP_cierre + B y par dinámico "
              "TD = C(θ)·ΔPef, con ΔPef = mín(ΔP, ΔP bloqueada) de cada caso y C interpolado en el ángulo de "
              "apertura. Se dimensiona con el mayor de ambos por un margen de 1,25.")
    else:
        doc.p("Método de fuerzas para válvulas de vástago deslizante: F = A + B + C + D, donde A es el "
              "desbalance estático del fluido sobre el obturador y el vástago, B la carga de asiento de la clase de "
              "cierre, C la fricción de la empaquetadura y D las fuerzas adicionales. Presiones manométricas a "
              "cierre en psi, áreas en in² y fuerzas en lbf (con su equivalente en N).")
    doc.tabla([[a, b, c] for a, b, c in ra.fuerzas], ["Componente", "Valor", "Fórmula / base"], anchos=[30, 25, 45])
    if ra.pasos:
        doc.tabla([[a, b, c] for a, b, c in ra.pasos], ["Paso", "Ecuación", "Resultado"], anchos=[28, 47, 25])
    doc.tabla(_filas_actuador(res, pr), ["Dato / resultado", "Valor"], anchos=[35, 65])
    doc.h3("Posicionador")
    doc.tabla([[f, v] for f, v, _ in eq.posicionador.filas], anchos=[35, 65])
    if eq.posicionador.bloqueadas:
        doc.lista([f"{s}: {m}." for s, m in eq.posicionador.bloqueadas.items()])
    if eq.accesorios.filas:
        doc.h3("Accesorios")
        doc.tabla([[f, v] for f, v, _ in eq.accesorios.filas], anchos=[35, 65])
    doc.nota("Tamaños de actuador de series genéricas (no corresponden a ningún fabricante). El fabricante confirma el "
             "modelo, el resorte y el ajuste de banco.")


def _seccion_equipo_informe(doc: Documento, res: Resultado, pr: Presentacion, n: int):
    eq = getattr(res, "equipo", None)
    doc.h2(f"{n}. Actuador, posicionador y accesorios")
    if eq is None:
        doc.p(_SIN_EQUIPO + ".")
        return
    ra = eq.actuador
    doc.p(("Actuador apto: " if ra.ok else "El actuador NO cumple: ") + (ra.seleccion or "—") + ".")
    doc.tabla([[a, b] for a, b, _ in ra.fuerzas], ["Componente", "Valor"], anchos=[45, 55])
    doc.tabla(_filas_actuador(res, pr), anchos=[35, 65])
    doc.h3("Posicionador")
    doc.tabla([[f, v] for f, v, _ in eq.posicionador.filas], anchos=[35, 65])
    doc.p(f"Acción sugerida: {eq.posicionador.accion_sugerida}. Set de aire sugerido: {eq.posicionador.set_sugerido}.")
    if eq.accesorios.filas:
        doc.h3("Accesorios")
        doc.tabla([[f, v] for f, v, _ in eq.accesorios.filas], anchos=[35, 65])
